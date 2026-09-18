# -*- coding: utf-8 -*-
"""
DIP分组费率查询器 - 数据预处理脚本
读取两个Excel源文件 + docx文档，生成精简的 .js 数据文件（file:// 协议下可用 script 标签加载）。
如源数据更新，在本机安装 python + openpyxl + pypinyin + python-docx 后重跑本脚本即可。
"""
import json
import re
import openpyxl
from pypinyin import lazy_pinyin, Style
import docx
from docx.document import Document as _Doc
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn

OUT_DIR = "DIP分组费率查询器"

# ---------- 1. ICD诊断库：诊断编码、诊断名称、拼音首字母 ----------
def make_initials(name):
    """生成检索用拼音首字母串：中文取首字母，字母/数字原样小写，其余忽略"""
    out = []
    for ch in name:
        if ch.isascii():
            if ch.isalnum():
                out.append(ch.lower())
            continue
        py = lazy_pinyin(ch, style=Style.FIRST_LETTER)
        if py and py[0]:
            out.append(py[0][0].lower())
    return "".join(out)

wb = openpyxl.load_workbook("_tmp_icd.xlsx", read_only=True)
ws = wb["湖北2.0版疾病完整库"]
diagnoses = []
seen = set()
for code, _ext, name in ws.iter_rows(min_row=2, values_only=True):
    if not code or not name:
        continue
    code = str(code).strip()
    name = str(name).strip()
    if code in seen and name in [d[1] for d in diagnoses if d[0] == code]:
        continue
    seen.add(code)
    diagnoses.append([code, name, make_initials(name)])
wb.close()
print("诊断条数:", len(diagnoses))

# ---------- 2. DIP目录 ----------
def prefix4(code):
    """去小数点后取前4位字母数字，大写。如 K81.000x001 -> K810"""
    s = re.sub(r"[^A-Za-z0-9]", "", str(code)).upper()
    return s[:4]

wb = openpyxl.load_workbook("_tmp_dip.xlsx", read_only=True)

# 核心病组：病种编码、病种名称、诊断编码、三级分值（按前4位建索引，分值降序）
ws = wb["核心病组"]
core = {}
for row in ws.iter_rows(min_row=2, values_only=True):
    bcode, bname, dcode, _avg, score = row[0], row[1], row[2], row[3], row[4]
    if not bcode or score is None:
        continue
    key = prefix4(dcode)
    core.setdefault(key, []).append([str(bcode).strip(), str(bname).strip(), round(float(score), 3)])
for k in core:
    core[k].sort(key=lambda r: -r[2])
print("核心病组前四位键数:", len(core), " 总条数:", sum(len(v) for v in core.values()))

# 综合病组：编码(诊断部分-治疗方式)、名称、三级分值（按诊断部分建索引，分值降序）
ws = wb["综合病组"]
comp = {}
for row in ws.iter_rows(min_row=2, values_only=True):
    _cls, code, name, _avg, score = row[0], row[1], row[2], row[3], row[4]
    if not code or score is None:
        continue
    code = str(code).strip()
    diag = code.split("-")[0].upper()
    comp.setdefault(diag, []).append([code, str(name).strip(), round(float(score), 3)])
for k in comp:
    comp[k].sort(key=lambda r: -r[2])
print("综合病组类目键数:", len(comp), " 总条数:", sum(len(v) for v in comp.values()))
wb.close()

# ---------- 3. docx -> 排版HTML ----------
def iter_block_items(parent):
    body = parent.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, parent)
        elif child.tag == qn("w:tbl"):
            yield Table(child, parent)

def esc(t):
    return (t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))

d = docx.Document("甲乳1(2).docx")
html = []
for block in iter_block_items(d):
    if isinstance(block, Paragraph):
        text = block.text.strip()
        if not text:
            continue
        style = block.style.name
        if style == "Heading 1":
            html.append(f"<h3 class='doc-h1'>{esc(text)}</h3>")
        elif style == "Heading 2":
            html.append(f"<h4 class='doc-h2'>{esc(text)}</h4>")
        else:
            # 费用数字高亮
            t = esc(text)
            t = re.sub(r"(\d[\d,]*\.\d{1,2})\s*(元)?",
                       r"<span class='fee'>\1\2</span>", t)
            html.append(f"<p>{t}</p>")
    else:
        rows_html = []
        for r in block.rows:
            cells = "".join(f"<td>{esc(c.text.strip())}</td>" for c in r.cells)
            rows_html.append(f"<tr>{cells}</tr>")
        if rows_html:
            # 首行作为表头
            rows_html[0] = rows_html[0].replace("<td>", "<th>").replace("</td>", "</th>")
            html.append("<table class='doc-table'>" + "".join(rows_html) + "</table>")
doc_html = "\n".join(html)
print("docx转换HTML长度:", len(doc_html))

# ---------- 4. 写出 .js 数据文件 ----------
import os
os.makedirs(f"{OUT_DIR}/data", exist_ok=True)

def dump_js(fname, varname, obj):
    with open(f"{OUT_DIR}/data/{fname}", "w", encoding="utf-8") as f:
        f.write(f"window.{varname}=")
        json.dump(obj, f, ensure_ascii=False, separators=(",", ":"))
        f.write(";")
    print("写出", fname, os.path.getsize(f"{OUT_DIR}/data/{fname}") // 1024, "KB")

dump_js("diagnoses.js", "DIAGNOSES", diagnoses)
dump_js("core.js", "DIP_CORE", core)
dump_js("comprehensive.js", "DIP_COMP", comp)
dump_js("doc.js", "DOC_HTML", doc_html)
print("完成")
